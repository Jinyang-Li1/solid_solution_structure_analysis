"""
disorder_analysis_with_diagrams.py - A tool for CSD refcode disorder analysis that combines
disorder detail extraction and Mercury-style structure diagram generation.
"""

import csv
import os
from io import BytesIO
from tqdm import tqdm
from ccdc.io import EntryReader
from ccdc.diagram import DiagramGenerator
from openpyxl import Workbook
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage


# Helper Functions
def pixels_to_points(pixels, dpi=96):
    """Convert pixels to Excel row height (in points)"""
    return round(pixels * (72 / dpi), 1)


# Function 1: Extract Disorder Details
def extract_disorder_details(input_csv, intermediate_csv):
    """Read refcodes from the input CSV, extract raw disorder details from the CSD,
    and write results to an intermediate CSV with "disorder_detail" as the 4th column."""
    try:
        # Read input CSV and prepare headers
        with open(input_csv, 'r', encoding='utf-8-sig') as f:
            reader = csv.reader(f)
            headers = next(reader)
            # Insert "disorder_detail" as the 4th column
            new_headers = headers[:3] + ["disorder_detail"] + headers[3:]
            # Filter empty rows to avoid processing invalid data
            data_rows = [row for row in reader if row]
            total_rows = len(data_rows)

            if not total_rows:
                print("No valid data in input CSV")
                return False

        # Extract disorder details and write to intermediate CSV
        with EntryReader('CSD') as entry_reader, \
                open(intermediate_csv, 'w', newline='', encoding='utf-8-sig') as f:

            writer = csv.writer(f)
            writer.writerow(new_headers)

            # Process each row with progress bar
            for row in tqdm(data_rows, total=total_rows, desc="Extracting Disorder Details"):
                refcode = row[0].strip() if row else ""
                disorder_detail = ""

                # Fetch disorder details if refcode is valid
                if refcode:
                    try:
                        disorder_detail = str(entry_reader.entry(refcode).disorder_details).strip()
                    except Exception as e:
                        # Truncate long error messages for readability
                        tqdm.write(f"Failed to extract {refcode}: {str(e)[:30]}...")

                # Write row with disorder details (preserves original columns)
                writer.writerow(row[:3] + [disorder_detail] + row[3:])

        print(f"Disorder details extracted → {intermediate_csv}")
        return True

    except Exception as e:
        print(f"Extraction failed: {str(e)}")
        return False


# Function 2: Generate Excel with Structure Diagrams
def generate_diagram_excel(intermediate_csv, output_excel):
    """GGenerate an Excel file with embedded Mercury-style structure diagrams from the intermediate CSV.
    Adds a "molecule_diagram" column (5th column) with images; adjusts column widths and row heights."""
    try:
        # Initialize diagram generator
        dg = DiagramGenerator()
        dg.settings.size = (350, 300)
        dg.settings.show_hydrogens = dg.settings.show_atom_labels = True
        dg.settings.atom_label_format = "{label}({occupancy:.2f})"

        # Configure Excel column widths (column number → width in characters)
        COLUMN_WIDTHS = {1: 12, 2: 15, 3: 15, 4: 40, 5: 30}

        # Read intermediate CSV and prepare data
        with open(intermediate_csv, 'r', encoding='utf-8-sig') as f:
            reader = csv.reader(f)
            headers = next(reader)
            # Insert "molecule_diagram" column after the 4th column
            headers.insert(4, "molecule_diagram")
            # Filter empty rows
            data_rows = [row for row in reader if row]
            total_rows = len(data_rows)

            if not total_rows:
                print("No valid data in intermediate CSV")
                return

        wb = Workbook()
        ws = wb.active
        ws.title = "CSD_Disorder_Analysis"

        # Write headers and set column widths
        for col_idx, header in enumerate(headers, 1):
            ws.cell(row=1, column=col_idx, value=header)
            if col_idx in COLUMN_WIDTHS:
                ws.column_dimensions[get_column_letter(col_idx)].width = COLUMN_WIDTHS[col_idx]

        # Track successful diagram generations
        success_count = 0

        # Generate diagrams and populate Excel rows
        with EntryReader('CSD') as entry_reader:
            for row_idx, csv_row in tqdm(enumerate(data_rows, 2), total=total_rows, desc="Generating Diagrams"):
                refcode = csv_row[0].strip()

                # Write data to first 4 columns (preserve intermediate CSV content)
                for col_idx in range(1, 5):
                    cell_value = csv_row[col_idx - 1] if len(csv_row) >= col_idx else ""
                    ws.cell(row=row_idx, column=col_idx, value=cell_value)

                # Generate and embed diagram if refcode is valid
                if refcode:
                    try:
                        # Generate diagram as PIL Image
                        diagram_pil = dg.image(entry_reader.entry(refcode))
                        img_buffer = BytesIO()
                        diagram_pil.save(img_buffer, format='PNG', dpi=(96, 96))
                        img_buffer.seek(0) # Reset buffer pointer to start

                        # Insert image into Excel and adjust row height
                        excel_image = OpenpyxlImage(img_buffer)
                        excel_image.width, excel_image.height = 350, 300
                        ws.add_image(excel_image, f'E{row_idx}')
                        ws.row_dimensions[row_idx].height = pixels_to_points(300)
                        success_count += 1
                    except Exception as e:
                        tqdm.write(f"Diagram generation failed for {refcode}: {str(e)[:30]}...")
                        ws.cell(row=row_idx, column=5, value="Generation Failed")
                        ws.row_dimensions[row_idx].height = 20
                else:
                    ws.cell(row=row_idx, column=5, value="No Refcode")
                    ws.row_dimensions[row_idx].height = 20

        # Create output directory if it doesn't exist (fix for path errors)
        dir_path = os.path.dirname(output_excel)
        if dir_path:
            os.makedirs(dir_path, exist_ok=True)
        # Save Excel file
        wb.save(output_excel)

        print(f"Excel file generated → {output_excel}")
        print(f"Statistics: Successfully generated {success_count}/{total_rows} diagrams")

    except Exception as e:
        print(f"Excel generation failed: {str(e)}")


def main():
    INPUT_CSV = "11. same_molecular_formula_in_same_atom_disorder.csv"
    INTERMEDIATE_CSV = "12. same_molecular_formula_with_disorder.csv"
    OUTPUT_EXCEL = "12. same_molecular_formula_with_diagrams.xlsx"

    print("Starting CSD Disorder Analysis & Diagram Generation...")
    if extract_disorder_details(INPUT_CSV, INTERMEDIATE_CSV):
        generate_diagram_excel(INTERMEDIATE_CSV, OUTPUT_EXCEL)
    print("Workflow completed")


if __name__ == "__main__":
    main()